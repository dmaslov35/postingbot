import vk_api
import time
import random
import json
import logging
import os
from datetime import datetime, timedelta
from openpyxl import load_workbook

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('logs/bot.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

POSTS_FILE = 'published_posts.json'
STATS_FILE = 'bot_stats.json'

MONTHS_RU = [
    'Января', 'Февраля', 'Марта', 'Апреля', 'Мая', 'Июня',
    'Июля', 'Августа', 'Сентября', 'Октября', 'Ноября', 'Декабря'
]

def load_config():
    with open('config.json', 'r', encoding='utf-8') as f:
        return json.load(f)

def load_settings():
    try:
        with open('settings.json', 'r', encoding='utf-8') as f:
            return json.load(f)
    except:
        return {"post_interval": 30, "check_interval": 7200, "delete_delay": 10}

def load_groups(excel_file):
    groups = []
    try:
        wb = load_workbook(excel_file)
        sheet = wb.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0]:
                groups.append(int(row[0]))
        logger.info(f"Loaded {len(groups)} groups from {excel_file}")
    except Exception as e:
        logger.error(f"Error reading Excel: {e}")
    return groups

def load_messages(messages_file):
    messages = []
    try:
        wb = load_workbook(messages_file)
        sheet = wb.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0]:
                messages.append(str(row[0]))
        logger.info(f"Loaded {len(messages)} messages from {messages_file}")
    except Exception as e:
        logger.error(f"Error reading messages: {e}")
    return messages

def load_published_posts():
    if os.path.exists(POSTS_FILE):
        with open(POSTS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def save_published_posts(posts):
    with open(POSTS_FILE, 'w', encoding='utf-8') as f:
        json.dump(posts, f, ensure_ascii=False, indent=2)

def load_stats():
    try:
        with open(STATS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except:
        return {'success': 0, 'errors': [], 'accounts_status': {}}

def save_stats(stats):
    with open(STATS_FILE, 'w', encoding='utf-8') as f:
        json.dump(stats, f, ensure_ascii=False, indent=2)

def check_pause():
    if os.path.exists('pause.flag'):
        logger.info("PAUSE command received from Telegram. Waiting...")
        while os.path.exists('pause.flag'):
            time.sleep(5)
        logger.info("RESUME command received. Continuing work!")

def smart_sleep(seconds):
    for _ in range(seconds):
        check_pause()
        time.sleep(1)

def check_post_status(vk, owner_id, post_id):
    try:
        response = vk.wall.getById(posts=f"{owner_id}_{post_id}")
        if response:
            return 'published'
        return 'not_found'
    except:
        return 'error'

def generate_dates():
    today = datetime.now()
    dates_comma = []
    for i in range(3):
        date = today + timedelta(days=i)
        dates_comma.append(str(date.day))
    month_name = MONTHS_RU[today.month - 1]
    dates_str_1 = ', '.join(dates_comma) + ' ' + month_name
    start_day = today.day
    end_day = (today + timedelta(days=2)).day
    dates_str_2 = f"{start_day}-{end_day} {month_name}"
    return random.choice([dates_str_1, dates_str_2])

def check_and_delete_old_posts():
    """Проверка и удаление старых постов"""
    logger.info("\n🔍 Checking and deleting old posts...")
    all_posts = load_published_posts()
    updated_posts = []
    
    settings = load_settings()
    delete_delay_seconds = settings['delete_delay'] * 60
    
    deleted_count = 0
    for post in all_posts:
        if post['status'] == 'deleted':
            updated_posts.append(post)
            continue
        
        try:
            vk_session = vk_api.VkApi(token=post['account_token'], api_version='5.199')
            vk = vk_session.get_api()
            
            status = check_post_status(vk, post['owner_id'], post['post_id'])
            
            if status == 'published':
                if 'published_on_wall_at' not in post:
                    post['published_on_wall_at'] = datetime.now().isoformat()
                
                published_time = datetime.fromisoformat(post['published_on_wall_at'])
                if datetime.now() - published_time > timedelta(seconds=delete_delay_seconds):
                    try:
                        vk.wall.delete(owner_id=post['owner_id'], post_id=post['post_id'])
                        post['status'] = 'deleted'
                        deleted_count += 1
                        logger.info(f"️ Deleted post {post['post_id']} from group {post['owner_id']}")
                    except Exception as e:
                        logger.error(f"Failed to delete post {post['post_id']}: {e}")
            
            updated_posts.append(post)
        except Exception as e:
            logger.error(f"Error checking post {post['post_id']}: {e}")
            updated_posts.append(post)
    
    save_published_posts(updated_posts)
    logger.info(f"✅ Post check complete. Deleted: {deleted_count}\n")

def main():
    config = load_config()
    settings = load_settings()
    accounts = [acc for acc in config['vk_accounts'] if acc.get('enabled', True)]
    groups = load_groups(config['settings']['excel_file'])
    messages = load_messages('messages.xlsx')

    if not accounts or not groups or not messages:
        logger.error("No accounts, groups or messages!")
        return

    logger.info(f"Starting bot. Accounts: {len(accounts)}, Groups: {len(groups)}, Messages: {len(messages)}")
    logger.info(f"Settings: Post interval {settings['post_interval']}s, Check every {settings['check_interval']}s, Delete after {settings['delete_delay']}m")

    acc_index = 0
    group_index = 0
    posts_since_last_check = 0
    
    while True:
        current_acc = accounts[acc_index % len(accounts)]
        current_group = groups[group_index % len(groups)]
        
        logger.info(f"\n{'='*60}")
        logger.info(f"Account {acc_index % len(accounts) + 1}/{len(accounts)}: {current_acc['user_id']} | Group: {current_group}")
        logger.info(f"{'='*60}")
        
        try:
            vk_session = vk_api.VkApi(token=current_acc['token'], api_version='5.199')
            vk = vk_session.get_api()
            logger.info("Authorization successful")
            
            check_pause()
            
            try:
                random_message = random.choice(messages)
                account_phone = current_acc.get('phone', 'Не указан')
                current_dates = generate_dates()
                
                final_message = random_message.replace('{PHONE}', account_phone)
                final_message = final_message.replace('{DATES}', current_dates)
                
                response = vk.wall.post(owner_id=current_group, message=final_message)
                post_id = response['post_id']
                
                all_posts = load_published_posts()
                all_posts.append({
                    'owner_id': current_group,
                    'post_id': post_id,
                    'created_at': datetime.now().isoformat(),
                    'status': 'pending',
                    'account_token': current_acc['token']
                })
                save_published_posts(all_posts)
                
                logger.info(f"✅ POSTED to group {current_group} (Post ID: {post_id})")
                logger.info(f"Phone: {account_phone} | Dates: {current_dates}")
                
                # Обновляем статистику
                all_stats = load_stats()
                all_stats['success'] += 1
                all_stats['accounts_status'][str(current_acc['user_id'])] = 'OK'
                save_stats(all_stats)
                
                posts_since_last_check += 1
                
            except vk_api.exceptions.ApiError as e:
                logger.warning(f"❌ Group {current_group}: Error {e.code}")
                
                # Обновляем статистику с ошибкой
                all_stats = load_stats()
                all_stats['errors'].append({
                    'group': current_group,
                    'error': str(e),
                    'account': current_acc['user_id']
                })
                all_stats['errors'] = all_stats['errors'][-50:]
                all_stats['accounts_status'][str(current_acc['user_id'])] = f'Error {e.code}'
                save_stats(all_stats)
                
            except Exception as e:
                logger.error(f"❌ System error {current_group}: {e}")
            
            # Переходим к следующему аккаунту и группе
            acc_index += 1
            group_index += 1
            
            smart_sleep(settings['post_interval'])
            
            # Проверка и удаление старых постов каждые 10 публикаций
            if posts_since_last_check >= 10:
                check_and_delete_old_posts()
                posts_since_last_check = 0
            
        except Exception as e:
            logger.error(f"Critical error: {e}")
            time.sleep(60)

if __name__ == '__main__':
    main()
