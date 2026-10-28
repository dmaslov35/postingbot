import asyncio
import html
import json
import logging
import os
import signal
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.error import BadRequest
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config.json"
SETTINGS_FILE = BASE_DIR / "settings.json"
STATS_FILE = BASE_DIR / "bot_stats.json"
STATE_FILE = BASE_DIR / "bot_state.json"
POSTS_FILE = BASE_DIR / "published_posts.json"
PAUSE_FILE = BASE_DIR / "pause.flag"
MESSAGES_FILE = BASE_DIR / "messages.xlsx"
GROUPS_FILE = BASE_DIR / "chats.xlsx"
LOG_FILE = BASE_DIR / "logs" / "bot.log"
ADMIN_LOG_FILE = BASE_DIR / "telegram.log"

DEFAULT_SETTINGS = {
    "post_interval": 30,
    "check_interval": 7200,
    "delete_delay": 10,
}

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler(ADMIN_LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger("telegram_admin")


def atomic_write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent), text=True)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        with path.open("r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        logger.error("Ошибка чтения %s: %s", path.name, e)
        return default


def load_config() -> Dict[str, Any]:
    data = load_json(CONFIG_FILE, {})
    if not isinstance(data, dict):
        return {}
    return data


CONFIG = load_config()
BOT_TOKEN = CONFIG.get("telegram", {}).get("bot_token", "")
telegram_cfg = CONFIG.get("telegram", {})
if "admin_ids" in telegram_cfg:
    ADMIN_IDS = [int(x) for x in telegram_cfg.get("admin_ids", [])]
elif "admin_id" in telegram_cfg:
    ADMIN_IDS = [int(telegram_cfg["admin_id"])]
else:
    ADMIN_IDS = []


def is_admin(update: Update) -> bool:
    return bool(update.effective_user and update.effective_user.id in ADMIN_IDS)


def load_settings() -> Dict[str, int]:
    raw = load_json(SETTINGS_FILE, DEFAULT_SETTINGS.copy())
    result = DEFAULT_SETTINGS.copy()
    if isinstance(raw, dict):
        for key, default in DEFAULT_SETTINGS.items():
            try:
                value = int(raw.get(key, default))
                if value > 0:
                    result[key] = value
            except (TypeError, ValueError):
                pass
    return result


def save_settings(settings: Dict[str, int]) -> None:
    atomic_write_json(SETTINGS_FILE, settings)


def load_stats() -> Dict[str, Any]:
    data = load_json(STATS_FILE, {})
    return data if isinstance(data, dict) else {}


def load_state() -> Dict[str, Any]:
    data = load_json(STATE_FILE, {})
    return data if isinstance(data, dict) else {}


def load_posts() -> List[Dict[str, Any]]:
    data = load_json(POSTS_FILE, [])
    return data if isinstance(data, list) else []


def load_groups() -> List[int]:
    groups = []
    try:
        wb = load_workbook(GROUPS_FILE, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                groups.append(int(row[0]))
        wb.close()
    except Exception:
        logger.exception("Ошибка чтения chats.xlsx")
    return groups


def load_messages() -> List[str]:
    messages = []
    try:
        wb = load_workbook(MESSAGES_FILE, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                messages.append(str(row[0]))
        wb.close()
    except Exception:
        logger.exception("Ошибка чтения messages.xlsx")
    return messages


def save_messages(messages: List[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Messages"
    ws.append(["message"])
    for msg in messages:
        ws.append([msg])
    tmp = MESSAGES_FILE.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, MESSAGES_FILE)


def bot_process_info() -> Dict[str, str]:
    try:
        r = subprocess.run(
            ["systemctl", "show", "vkbot.service", "--property=ActiveState,SubState,MainPID"],
            capture_output=True, text=True, timeout=5, check=False
        )
        info = {}
        for line in r.stdout.splitlines():
            if "=" in line:
                k, v = line.split("=", 1)
                info[k] = v
        return info
    except Exception as e:
        logger.warning("Не удалось получить systemd status: %s", e)
        return {}


def bot_is_active() -> bool:
    return bot_process_info().get("ActiveState") == "active"


def restart_bot_service() -> tuple[bool, str]:
    try:
        r = subprocess.run(
            ["systemctl", "restart", "vkbot.service"],
            capture_output=True, text=True, timeout=20, check=False
        )
        if r.returncode != 0:
            return False, (r.stderr or r.stdout or "systemctl error").strip()
        return bot_is_active(), "Сервис перезапущен"
    except Exception as e:
        return False, str(e)


def status_label(raw: Any) -> tuple[str, str]:
    if isinstance(raw, dict):
        value = str(raw.get("status", "UNKNOWN"))
    else:
        value = str(raw)

    if value == "OK":
        return "🟢", "Работает"
    if value.startswith("AUTH_ERROR"):
        return "🔴", value
    if value.startswith("TEMP_ERROR"):
        return "🟡", value
    if value == "CONFIG_ERROR":
        return "🔴", "Ошибка config"
    if value.startswith("VK_ERROR"):
        return "🟠", value
    return "⚪", value


def account_summary() -> tuple[str, int, int, int]:
    cfg = load_config()
    accounts = [a for a in cfg.get("vk_accounts", []) if a.get("enabled", True)]
    stats = load_stats()
    statuses = stats.get("accounts_status", {})
    state = load_state()
    cooldowns = state.get("account_cooldowns", {})

    ok = bad = unknown = 0
    lines = []
    now = datetime.now()

    for i, account in enumerate(accounts, 1):
        uid = str(account.get("user_id", f"#{i}"))
        raw = statuses.get(uid)
        emoji, label = status_label(raw if raw is not None else "Нет данных")

        cooldown_until = cooldowns.get(uid)
        cooldown_text = ""
        if cooldown_until:
            try:
                until = datetime.fromisoformat(cooldown_until)
                if until > now:
                    cooldown_text = f" · cooldown до {until.strftime('%H:%M')}"
            except (TypeError, ValueError):
                pass

        if label == "Работает":
            ok += 1
        elif raw is None:
            unknown += 1
        else:
            bad += 1

        lines.append(f"{emoji} {i}. {uid}: {label}{cooldown_text}")

    return "\n".join(lines), ok, bad, unknown


def main_dashboard_text() -> str:
    settings = load_settings()
    stats = load_stats()
    state = load_state()
    posts = load_posts()
    cfg = load_config()
    groups = load_groups()
    messages = load_messages()

    enabled_accounts = [a for a in cfg.get("vk_accounts", []) if a.get("enabled", True)]
    pending = sum(1 for p in posts if p.get("status") != "deleted")
    errors = stats.get("errors", [])
    paused = PAUSE_FILE.exists()
    active = bot_is_active()

    account_index = int(state.get("account_index", 0)) if enabled_accounts else 0
    group_index = int(state.get("group_index", 0)) if groups else 0

    service = "🟢 запущен" if active else "🔴 остановлен"
    posting = "⏸ на паузе" if paused else "▶️ работает"

    return (
        "🤖 VK BOT — панель управления\n\n"
        f"⚙️ Сервис: {service}\n"
        f"📤 Публикация: {posting}\n\n"
        f"📊 Всего успешных: {int(stats.get('success', 0))}\n"
        f"❌ Ошибок в истории: {len(errors)}\n"
        f"🗑 В очереди удаления: {pending}\n\n"
        f"👥 Аккаунтов: {len(enabled_accounts)}\n"
        f"📋 Групп: {len(groups)}\n"
        f"📝 Текстов: {len(messages)}\n\n"
        f"🔄 Позиция: аккаунт {account_index + 1 if enabled_accounts else 0}/{len(enabled_accounts)}, "
        f"группа {group_index + 1 if groups else 0}/{len(groups)}\n\n"
        f"⏱ Между постами: {settings['post_interval']} сек\n"
        f"🔎 Проверка: {settings['check_interval']} сек\n"
        f"🗑 Удаление: через {settings['delete_delay']} мин"
    )


def main_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📊 Статистика", callback_data="stats"),
         InlineKeyboardButton("👥 Аккаунты", callback_data="accounts")],
        [InlineKeyboardButton("📝 Тексты", callback_data="messages_menu"),
         InlineKeyboardButton("📋 Группы", callback_data="groups_page_0")],
        [InlineKeyboardButton("⏱ Настройки", callback_data="settings_menu"),
         InlineKeyboardButton("📈 Логи", callback_data="logs")],
        [InlineKeyboardButton("⏸ Пауза", callback_data="pause"),
         InlineKeyboardButton("▶️ Продолжить", callback_data="resume")],
        [InlineKeyboardButton("🔄 Обновить", callback_data="main_menu"),
         InlineKeyboardButton("♻️ Перезапуск", callback_data="restart_confirm")],
    ])


async def safe_edit(query, text: str, reply_markup=None, parse_mode=None) -> None:
    try:
        await query.message.edit_text(
            text=text, reply_markup=reply_markup, parse_mode=parse_mode,
            disable_web_page_preview=True
        )
    except BadRequest as e:
        if "Message is not modified" not in str(e):
            raise


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        if update.message:
            await update.message.reply_text("❌ Доступ запрещён")
        return
    context.user_data.pop("action", None)
    await update.message.reply_text(main_dashboard_text(), reply_markup=main_keyboard())


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        return
    context.user_data.pop("action", None)
    await update.message.reply_text("✅ Действие отменено.", reply_markup=main_keyboard())


async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if not is_admin(update):
        return

    data = query.data

    if data == "main_menu":
        context.user_data.pop("action", None)
        await safe_edit(query, main_dashboard_text(), main_keyboard())

    elif data == "stats":
        stats = load_stats()
        errors = stats.get("errors", [])
        posts = load_posts()
        pending = sum(1 for p in posts if p.get("status") != "deleted")
        text = (
            "📊 СТАТИСТИКА\n\n"
            f"✅ Успешных публикаций: {int(stats.get('success', 0))}\n"
            f"❌ Ошибок в истории: {len(errors)}\n"
            f"🗑 В очереди удаления: {pending}\n"
        )
        if errors:
            text += "\nПоследние 5 ошибок:\n"
            for err in errors[-5:]:
                code = err.get("code")
                code_text = f" [{code}]" if code is not None else ""
                text += (
                    f"• {err.get('time', '')} · acc {err.get('account')} · "
                    f"group {err.get('group')} · {err.get('type', '')}{code_text}\n"
                )
        await safe_edit(query, text, InlineKeyboardMarkup([
            [InlineKeyboardButton("🔄 Обновить", callback_data="stats")],
            [InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]
        ]))

    elif data == "accounts":
        lines, ok, bad, unknown = account_summary()
        text = (
            "👥 АККАУНТЫ\n\n"
            f"🟢 Работают по последнему статусу: {ok}\n"
            f"🔴/🟡 С ошибками: {bad}\n"
            f"⚪ Без данных: {unknown}\n\n"
            f"{lines or 'Нет аккаунтов'}"
        )
        await safe_edit(query, text[:4000], InlineKeyboardMarkup([
            [InlineKeyboardButton("🔄 Обновить", callback_data="accounts")],
            [InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]
        ]))

    elif data == "settings_menu":
        s = load_settings()
        text = (
            "⏱ НАСТРОЙКИ\n\n"
            f"Пауза между постами: {s['post_interval']} сек\n"
            f"Интервал проверки: {s['check_interval']} сек\n"
            f"Удаление поста: {s['delete_delay']} мин\n\n"
            "Изменения подхватываются ядром без перезапуска."
        )
        await safe_edit(query, text, InlineKeyboardMarkup([
            [InlineKeyboardButton("Пауза между постами", callback_data="set_post_interval")],
            [InlineKeyboardButton("Интервал проверки", callback_data="set_check_interval")],
            [InlineKeyboardButton("Время удаления", callback_data="set_delete_delay")],
            [InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]
        ]))

    elif data in {"set_post_interval", "set_check_interval", "set_delete_delay"}:
        context.user_data["action"] = data
        prompts = {
            "set_post_interval": "Введите паузу между постами в секундах (минимум 1):",
            "set_check_interval": "Введите интервал проверки в секундах (минимум 1):",
            "set_delete_delay": "Введите время до удаления поста в минутах (минимум 1):",
        }
        await safe_edit(query, prompts[data] + "\n\nОтмена: /cancel",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Отмена", callback_data="settings_menu")]]))

    elif data == "messages_menu":
        messages = load_messages()
        await safe_edit(query, f"📝 ТЕКСТЫ\n\nВсего текстов: {len(messages)}", InlineKeyboardMarkup([
            [InlineKeyboardButton("📋 Показать", callback_data="messages_page_0")],
            [InlineKeyboardButton("➕ Добавить", callback_data="add_message")],
            [InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]
        ]))

    elif data.startswith("messages_page_"):
        page = max(0, int(data.rsplit("_", 1)[1]))
        messages = load_messages()
        per_page = 8
        total_pages = max(1, (len(messages) + per_page - 1) // per_page)
        page = min(page, total_pages - 1)
        start_i = page * per_page
        page_items = messages[start_i:start_i + per_page]
        text = f"📝 Тексты — {page + 1}/{total_pages}\n\n"
        keyboard = []
        for idx, msg in enumerate(page_items, start_i + 1):
            preview = " ".join(msg.split())
            preview = preview[:45] + ("…" if len(preview) > 45 else "")
            keyboard.append([InlineKeyboardButton(f"{idx}. {preview}", callback_data=f"view_message_{idx}")])
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton("⬅️", callback_data=f"messages_page_{page-1}"))
        if page + 1 < total_pages:
            nav.append(InlineKeyboardButton("➡️", callback_data=f"messages_page_{page+1}"))
        if nav:
            keyboard.append(nav)
        keyboard.append([InlineKeyboardButton("◀️ Назад", callback_data="messages_menu")])
        await safe_edit(query, text, InlineKeyboardMarkup(keyboard))

    elif data.startswith("view_message_"):
        n = int(data.rsplit("_", 1)[1])
        messages = load_messages()
        if not 1 <= n <= len(messages):
            await safe_edit(query, "❌ Текст не найден.", InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Назад", callback_data="messages_menu")]]))
        else:
            msg = messages[n - 1]
            await safe_edit(query, f"📝 Текст #{n}\n\n{msg}"[:4000], InlineKeyboardMarkup([
                [InlineKeyboardButton("✏️ Изменить", callback_data=f"edit_message_{n}"),
                 InlineKeyboardButton("🗑 Удалить", callback_data=f"delete_confirm_{n}")],
                [InlineKeyboardButton("◀️ Назад", callback_data="messages_page_0")]
            ]))

    elif data == "add_message":
        context.user_data["action"] = "add_message"
        await safe_edit(query, "Отправьте новый текст.\n\nОтмена: /cancel",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Отмена", callback_data="messages_menu")]]))

    elif data.startswith("edit_message_"):
        n = int(data.rsplit("_", 1)[1])
        context.user_data["action"] = f"edit_message_{n}"
        await safe_edit(query, f"Отправьте новый вариант текста #{n}.\n\nОтмена: /cancel",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Отмена", callback_data=f"view_message_{n}")]]))

    elif data.startswith("delete_confirm_"):
        n = int(data.rsplit("_", 1)[1])
        await safe_edit(query, f"Удалить текст #{n}?", InlineKeyboardMarkup([
            [InlineKeyboardButton("✅ Да, удалить", callback_data=f"delete_message_{n}")],
            [InlineKeyboardButton("❌ Нет", callback_data=f"view_message_{n}")]
        ]))

    elif data.startswith("delete_message_"):
        n = int(data.rsplit("_", 1)[1])
        messages = load_messages()
        if 1 <= n <= len(messages):
            messages.pop(n - 1)
            save_messages(messages)
            text = f"✅ Текст #{n} удалён."
        else:
            text = "❌ Текст не найден."
        await safe_edit(query, text, InlineKeyboardMarkup([[InlineKeyboardButton("◀️ К текстам", callback_data="messages_menu")]]))

    elif data.startswith("groups_page_"):
        page = max(0, int(data.rsplit("_", 1)[1]))
        groups = load_groups()
        per_page = 15
        total_pages = max(1, (len(groups) + per_page - 1) // per_page)
        page = min(page, total_pages - 1)
        start_i = page * per_page
        page_items = groups[start_i:start_i + per_page]
        keyboard = []
        for idx, gid in enumerate(page_items, start_i + 1):
            keyboard.append([InlineKeyboardButton(f"{idx}. {gid}", url=f"https://vk.com/club{abs(gid)}")])
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton("⬅️", callback_data=f"groups_page_{page-1}"))
        if page + 1 < total_pages:
            nav.append(InlineKeyboardButton("➡️", callback_data=f"groups_page_{page+1}"))
        if nav:
            keyboard.append(nav)
        keyboard.append([InlineKeyboardButton("◀️ Назад", callback_data="main_menu")])
        await safe_edit(query, f"📋 ГРУППЫ — {page + 1}/{total_pages}\nВсего: {len(groups)}", InlineKeyboardMarkup(keyboard))

    elif data == "logs":
        try:
            lines = LOG_FILE.read_text(encoding="utf-8", errors="replace").splitlines()[-15:]
            log_text = "\n".join(lines)
            # Telegram HTML: обязательно экранируем содержимое логов.
            text = "📈 Последние строки лога:\n\n<pre>" + html.escape(log_text[-3300:]) + "</pre>"
        except OSError:
            text = "Логов пока нет."
        await safe_edit(query, text, InlineKeyboardMarkup([
            [InlineKeyboardButton("🔄 Обновить", callback_data="logs")],
            [InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]
        ]), parse_mode="HTML")

    elif data == "pause":
        PAUSE_FILE.touch(exist_ok=True)
        await safe_edit(query, "⏸ Публикации поставлены на паузу.\n\nУдаление старых постов продолжает работать.",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]]))

    elif data == "resume":
        try:
            PAUSE_FILE.unlink()
        except FileNotFoundError:
            pass
        await safe_edit(query, "▶️ Публикации возобновлены.",
                        InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]]))

    elif data == "restart_confirm":
        await safe_edit(query, "♻️ Перезапустить основной VK-бот?\n\nTelegram-админка останется работать.", InlineKeyboardMarkup([
            [InlineKeyboardButton("✅ Перезапустить", callback_data="restart_execute")],
            [InlineKeyboardButton("❌ Отмена", callback_data="main_menu")]
        ]))

    elif data == "restart_execute":
        ok, detail = await asyncio.to_thread(restart_bot_service)
        if ok:
            text = "✅ Основной VK-бот перезапущен и снова active."
        else:
            text = f"❌ Не удалось перезапустить VK-бот:\n{detail[:1000]}"
        await safe_edit(query, text, InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Назад", callback_data="main_menu")]]))


async def message_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update) or not update.message or not update.message.text:
        return

    action = context.user_data.get("action")
    if not action:
        return

    text = update.message.text

    if action == "add_message":
        messages = load_messages()
        messages.append(text)
        save_messages(messages)
        context.user_data.pop("action", None)
        await update.message.reply_text(f"✅ Текст добавлен. Всего: {len(messages)}", reply_markup=main_keyboard())
        return

    if action.startswith("edit_message_"):
        n = int(action.rsplit("_", 1)[1])
        messages = load_messages()
        if 1 <= n <= len(messages):
            messages[n - 1] = text
            save_messages(messages)
            answer = f"✅ Текст #{n} обновлён."
        else:
            answer = "❌ Текст не найден."
        context.user_data.pop("action", None)
        await update.message.reply_text(answer, reply_markup=main_keyboard())
        return

    if action in {"set_post_interval", "set_check_interval", "set_delete_delay"}:
        try:
            value = int(text.strip())
            if value < 1:
                raise ValueError
        except ValueError:
            await update.message.reply_text("❌ Введите целое число больше 0. Действие не отменено.")
            return

        settings = load_settings()
        key = action.removeprefix("set_")
        settings[key] = value
        save_settings(settings)
        context.user_data.pop("action", None)

        units = {
            "post_interval": "сек",
            "check_interval": "сек",
            "delete_delay": "мин",
        }
        await update.message.reply_text(
            f"✅ {key} = {value} {units[key]}\nИзменение применяется ядром автоматически.",
            reply_markup=main_keyboard()
        )


async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
    logger.exception("Ошибка Telegram-админки", exc_info=context.error)


def main() -> None:
    if not BOT_TOKEN:
        raise RuntimeError("В config.json отсутствует telegram.bot_token")
    if not ADMIN_IDS:
        raise RuntimeError("В config.json отсутствуют telegram.admin_id/admin_ids")

    os.chdir(BASE_DIR)

    logger.info("Telegram Admin v2 запущен. Администраторы: %s", ADMIN_IDS)

    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("cancel", cancel))
    app.add_handler(CallbackQueryHandler(button_handler))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, message_handler))
    app.add_error_handler(error_handler)
    app.run_polling()


if __name__ == "__main__":
    main()
